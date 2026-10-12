"""
report_generator.py
Generate Excel report dengan dua sheet:
  - DATA TERBARU  : snapshot hasil cek terkini
  - RIWAYAT CEK   : histori setiap run (append, tidak pernah dihapus)

Struktur kolom P1–P14:
  Setiap pertemuan memiliki 3 sub-kolom: PPT | Video | Quiz
  Isi: ✅ (ada) atau ❌ (belum ada)
  Jika course belum ditemukan / belum KRS → kolom kosong
"""

from __future__ import annotations

import logging
from datetime import datetime
from pathlib import Path
from typing import Any

import openpyxl
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from src.content_checker import MAX_PERTEMUAN, MIN_PERTEMUAN
from src import config

logger = logging.getLogger(__name__)

# ── Checklist symbols ─────────────────────────────────────────────────────────
CHECKLIST_YES = "✅"
CHECKLIST_NO  = "❌"

# ── Warna ─────────────────────────────────────────────────────────────────────
_COLORS = {
    "LENGKAP":                {"bg": "C6EFCE", "font": "276221"},
    "BELUM LENGKAP":          {"bg": "FFEB9C", "font": "9C6500"},
    "COURSE TIDAK DITEMUKAN": {"bg": "FFC7CE", "font": "9C0006"},
    "BELUM KRS":              {"bg": "D9D9D9", "font": "595959"},
    "API ERROR":              {"bg": "F4CCCC", "font": "CC0000"},
    "ERROR":                  {"bg": "F4CCCC", "font": "CC0000"},
    "HEADER_MAIN":            {"bg": "1F4E79", "font": "FFFFFF"},  # biru tua — info kelas
    "HEADER_PPT":             {"bg": "2E75B6", "font": "FFFFFF"},  # biru — PPT
    "HEADER_VIDEO":           {"bg": "375623", "font": "FFFFFF"},  # hijau tua — Video
    "HEADER_QUIZ":            {"bg": "7B3F00", "font": "FFFFFF"},  # coklat — Quiz
    "CELL_YES":               {"bg": "E2EFDA", "font": "276221"},  # hijau muda ✅
    "CELL_NO":                {"bg": "FCE4D6", "font": "C00000"},  # merah muda ❌
    "CELL_EMPTY":             {"bg": "F2F2F2", "font": "808080"},  # abu — tidak applicable
}

_thin = Side(style="thin", color="D0D0D0")
_border = Border(left=_thin, right=_thin, top=_thin, bottom=_thin)


def _fill(hex_color: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_color)


def _font(hex_color: str, bold: bool = False, size: int = 9) -> Font:
    return Font(color=hex_color, bold=bold, size=size)


def _style_header(cell, text: str, color_key: str = "HEADER_MAIN") -> None:
    c = _COLORS[color_key]
    cell.value = text
    cell.font = _font(c["font"], bold=True, size=9)
    cell.fill = _fill(c["bg"])
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = _border


def _style_status(cell, status: str) -> None:
    c = _COLORS.get(status, {"bg": "FFFFFF", "font": "000000"})
    cell.value = status
    cell.fill = _fill(c["bg"])
    cell.font = _font(c["font"], bold=True, size=9)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _border


def _style_data(cell, value: Any, wrap: bool = False, center: bool = False) -> None:
    cell.value = value
    cell.font = _font("000000", size=9)
    cell.alignment = Alignment(
        horizontal="center" if center else "left",
        vertical="center",
        wrap_text=wrap,
    )
    cell.border = _border


def _style_checklist(cell, has_content: bool | None) -> None:
    """
    Render sel checklist:
      True  → ✅ hijau muda
      False → ❌ merah muda
      None  → kosong / abu (belum KRS / not found)
    """
    if has_content is None:
        cell.value = ""
        cell.fill = _fill(_COLORS["CELL_EMPTY"]["bg"])
        cell.font = _font(_COLORS["CELL_EMPTY"]["font"], size=10)
    elif has_content:
        cell.value = CHECKLIST_YES
        cell.fill = _fill(_COLORS["CELL_YES"]["bg"])
        cell.font = _font(_COLORS["CELL_YES"]["font"], bold=True, size=10)
    else:
        cell.value = CHECKLIST_NO
        cell.fill = _fill(_COLORS["CELL_NO"]["bg"])
        cell.font = _font(_COLORS["CELL_NO"]["font"], bold=True, size=10)
    cell.alignment = Alignment(horizontal="center", vertical="center")
    cell.border = _border


def _auto_width(ws, col_widths: dict[int, float] | None = None,
                min_w: int = 5, max_w: int = 40) -> None:
    """Set lebar kolom otomatis, dengan override manual jika perlu."""
    for col in ws.columns:
        col_letter = get_column_letter(col[0].column)
        col_idx = col[0].column

        if col_widths and col_idx in col_widths:
            ws.column_dimensions[col_letter].width = col_widths[col_idx]
            continue

        max_len = min_w
        for cell in col:
            if cell.value:
                text = str(cell.value)
                # Emoji butuh lebih lebar
                char_len = sum(2 if ord(c) > 0x2000 else 1 for c in text)
                max_len = max(max_len, min(char_len, max_w))
        ws.column_dimensions[col_letter].width = max_len + 1


# ── Header layout ─────────────────────────────────────────────────────────────

# Kolom info kelas (sebelum P1)
_INFO_HEADERS = [
    "No", "Kelas", "Matakuliah", "SKS", "SMT",
    "Jadwal", "Dosen Pengajar", "Prodi", "Ruangan", "Kampus",
    "Status", "Status KRS", "Dosen Sesuai", "Moodle Course", "Course ID",
]
_N_INFO = len(_INFO_HEADERS)

# Kolom setelah pertemuan
_TAIL_HEADERS = ["Jadwal Quiz P1", "Status LMS", "Detail Kekurangan", "Last Check"]

# Lebar tetap untuk kolom info
_INFO_WIDTHS: dict[str, float] = {
    "No": 4, "Kelas": 8, "Matakuliah": 32, "SKS": 4, "SMT": 4,
    "Jadwal": 18, "Dosen Pengajar": 28, "Prodi": 26,
    "Ruangan": 7, "Kampus": 14, "Status": 8, "Status KRS": 13,
    "Dosen Sesuai": 7, "Moodle Course": 30, "Course ID": 9,
}

# Lebar kolom checklist (sempit karena hanya ✅/❌)
_CHECKLIST_COL_WIDTH = 5.0
_PERTEMUAN_HEADER_WIDTH = 6.0


def _write_headers(ws) -> dict[str, int]:
    """
    Tulis dua baris header:
      Row 1: info kelas | P1 (span 3) | P2 (span 3) | ... | P14 (span 3) | tail
      Row 2: (kosong)   | PPT Video Quiz | PPT Video Quiz | ...             | (kosong)

    Return: mapping nama → col index (1-based) untuk kolom tail
    """
    # Row 1 — info kelas
    for col_idx, header in enumerate(_INFO_HEADERS, start=1):
        cell = ws.cell(row=1, column=col_idx)
        _style_header(cell, header, "HEADER_MAIN")
        ws.merge_cells(
            start_row=1, start_column=col_idx,
            end_row=2, end_column=col_idx,
        )

    # Row 1 — P1..P14 group header (merge 3 kolom per pertemuan)
    col_start = _N_INFO + 1
    for p in range(MIN_PERTEMUAN, MAX_PERTEMUAN + 1):
        cell = ws.cell(row=1, column=col_start)
        _style_header(cell, f"P{p}", "HEADER_MAIN")
        ws.merge_cells(
            start_row=1, start_column=col_start,
            end_row=1, end_column=col_start + 2,
        )
        # Row 2 — sub-header PPT / Video / Quiz
        _style_header(ws.cell(row=2, column=col_start),     "PPT",   "HEADER_PPT")
        _style_header(ws.cell(row=2, column=col_start + 1), "Video", "HEADER_VIDEO")
        _style_header(ws.cell(row=2, column=col_start + 2), "Quiz",  "HEADER_QUIZ")
        col_start += 3

    # Row 1 — tail headers
    tail_col_map: dict[str, int] = {}
    for header in _TAIL_HEADERS:
        cell = ws.cell(row=1, column=col_start)
        _style_header(cell, header, "HEADER_MAIN")
        ws.merge_cells(
            start_row=1, start_column=col_start,
            end_row=2, end_column=col_start,
        )
        tail_col_map[header] = col_start
        col_start += 1

    return tail_col_map


def _write_data_terbaru(ws, results: list[dict]) -> None:
    ws.title = "DATA TERBARU"

    tail_col_map = _write_headers(ws)

    status_col    = tail_col_map["Status LMS"]
    detail_col    = tail_col_map["Detail Kekurangan"]
    lastchk_col   = tail_col_map["Last Check"]
    quizjdwl_col  = tail_col_map["Jadwal Quiz P1"]

    for row_idx, item in enumerate(results, start=3):
        meetings = item.get("meetings", {})
        lms_status = item.get("lms_status", item.get("match_status", ""))
        applicable = bool(meetings)  # False = belum KRS atau not found

        # ── Kolom info kelas ──────────────────────────────────────────────────
        info_values = [
            row_idx - 2,
            item.get("kode_kelas", ""),
            item.get("matakuliah", ""),
            item.get("sks", ""),
            item.get("smt", ""),
            item.get("jadwal", ""),
            item.get("dosen", ""),
            item.get("prodi", ""),
            item.get("ruangan_kelas", ""),
            item.get("kampus", ""),
            item.get("status_kelas", ""),
            item.get("match_status", ""),
            item.get("dosen_sesuai", "N/A"),  # ✅ / ❌ / N/A
            item.get("course_fullname", "") or item.get("moodle_shortname", ""),
            item.get("course_id", ""),
        ]
        for col_idx, val in enumerate(info_values, start=1):
            center = col_idx in (1, 4, 5, 9, 11, 13)  # No, SKS, SMT, Ruangan, Status, Dosen Sesuai
            _style_data(ws.cell(row=row_idx, column=col_idx), val, center=center)

        # Style "Status KRS" kolom 12 (setelah penambahan Status di kolom 11)
        _style_status(ws.cell(row=row_idx, column=12), item.get("match_status", ""))

        # ── Kolom checklist P1–P14 ─────────────────────────────────────────────
        col_start = _N_INFO + 1
        for p in range(MIN_PERTEMUAN, MAX_PERTEMUAN + 1):
            meeting = meetings.get(p)
            if not applicable or meeting is None:
                ppt_val = vid_val = quz_val = None
            else:
                ppt_val = meeting.has_ppt
                vid_val = meeting.has_video
                quz_val = meeting.has_quiz

            _style_checklist(ws.cell(row=row_idx, column=col_start),     ppt_val)
            _style_checklist(ws.cell(row=row_idx, column=col_start + 1), vid_val)
            _style_checklist(ws.cell(row=row_idx, column=col_start + 2), quz_val)
            col_start += 3

        # ── Tail columns ──────────────────────────────────────────────────────
        # Jadwal Quiz P1
        quiz_valid = item.get("quiz_jadwal_p1")  # True / False / None
        cell_qjdwl = ws.cell(row=row_idx, column=quizjdwl_col)
        if quiz_valid is True:
            cell_qjdwl.value = CHECKLIST_YES
            cell_qjdwl.fill = _fill(_COLORS["CELL_YES"]["bg"])
            cell_qjdwl.font = _font(_COLORS["CELL_YES"]["font"], bold=True, size=10)
        elif quiz_valid is False:
            cell_qjdwl.value = CHECKLIST_NO
            cell_qjdwl.fill = _fill(_COLORS["CELL_NO"]["bg"])
            cell_qjdwl.font = _font(_COLORS["CELL_NO"]["font"], bold=True, size=10)
        else:
            cell_qjdwl.value = "⚠"
            cell_qjdwl.fill = _fill("FFF2CC")  # kuning muda — tidak dicek
            cell_qjdwl.font = _font("7F6000", size=10)
        cell_qjdwl.alignment = Alignment(horizontal="center", vertical="center")
        cell_qjdwl.border = _border

        _style_status(ws.cell(row=row_idx, column=status_col), lms_status)
        _style_data(ws.cell(row=row_idx, column=detail_col),
                    item.get("detail_kekurangan", ""), wrap=True)
        _style_data(ws.cell(row=row_idx, column=lastchk_col),
                    item.get("last_check", ""), center=True)

    # ── Freeze 2 baris header + autofilter ───────────────────────────────────
    ws.freeze_panes = "A3"
    total_cols = _N_INFO + (MAX_PERTEMUAN - MIN_PERTEMUAN + 1) * 3 + len(_TAIL_HEADERS)
    ws.auto_filter.ref = (
        f"A2:{get_column_letter(total_cols)}{max(ws.max_row, 3)}"
    )

    # Row heights
    ws.row_dimensions[1].height = 22
    ws.row_dimensions[2].height = 22
    for r in range(3, len(results) + 3):
        ws.row_dimensions[r].height = 16

    # ── Lebar kolom ───────────────────────────────────────────────────────────
    # Info columns
    for col_idx, header in enumerate(_INFO_HEADERS, start=1):
        w = _INFO_WIDTHS.get(header, 12)
        ws.column_dimensions[get_column_letter(col_idx)].width = w

    # Checklist columns (sempit)
    col_start = _N_INFO + 1
    for p in range(MIN_PERTEMUAN, MAX_PERTEMUAN + 1):
        for offset in range(3):
            ws.column_dimensions[get_column_letter(col_start + offset)].width = _CHECKLIST_COL_WIDTH
        col_start += 3

    # Tail columns
    ws.column_dimensions[get_column_letter(detail_col)].width = 42
    ws.column_dimensions[get_column_letter(status_col)].width = 18
    ws.column_dimensions[get_column_letter(lastchk_col)].width = 16
    ws.column_dimensions[get_column_letter(quizjdwl_col)].width = 14


# ── RIWAYAT CEK ───────────────────────────────────────────────────────────────

_RIWAYAT_HEADERS = [
    "Timestamp", "Kelas", "Matakuliah", "Dosen",
    "Course ID", "Status LMS", "Detail Kekurangan",
]


def _append_riwayat(ws, results: list[dict], timestamp: str) -> None:
    """Append baris ke RIWAYAT CEK — tidak pernah menghapus histori."""
    if ws.max_row == 0 or (ws.max_row == 1 and ws.cell(1, 1).value is None):
        for col, h in enumerate(_RIWAYAT_HEADERS, start=1):
            _style_header(ws.cell(row=1, column=col), h)
        ws.freeze_panes = "A2"
        ws.auto_filter.ref = f"A1:{get_column_letter(len(_RIWAYAT_HEADERS))}1"

    status_col = _RIWAYAT_HEADERS.index("Status LMS") + 1
    detail_col = len(_RIWAYAT_HEADERS)

    for item in results:
        lms_status = item.get("lms_status", item.get("match_status", ""))
        if not lms_status:
            continue
        row_data = [
            timestamp,
            item.get("kode_kelas", ""),
            item.get("matakuliah", ""),
            item.get("dosen", ""),
            item.get("course_id", ""),
            lms_status,
            item.get("detail_kekurangan", ""),
        ]
        ws.append(row_data)
        last_row = ws.max_row
        for col in range(1, len(_RIWAYAT_HEADERS) + 1):
            cell = ws.cell(row=last_row, column=col)
            if col == status_col:
                _style_status(cell, str(row_data[col - 1]))
            else:
                _style_data(cell, row_data[col - 1], wrap=(col == detail_col))

    # Lebar kolom riwayat
    riwayat_widths = {1: 18, 2: 8, 3: 32, 4: 28, 5: 9, 6: 18, 7: 50}
    for col_idx, width in riwayat_widths.items():
        ws.column_dimensions[get_column_letter(col_idx)].width = width


# ── Public API ────────────────────────────────────────────────────────────────

def generate_report(
    results: list[dict],
    output_path: str | Path | None = None,
) -> Path:
    """
    Generate file Excel report.

    Args:
        results     : list hasil pengecekan (dari orchestrator)
        output_path : path output, default = output/report_YYYYMMDD_HHMMSS.xlsx

    Return: Path ke file yang dibuat/diupdate
    """
    output_dir = Path(config.OUTPUT_DIR)
    output_dir.mkdir(parents=True, exist_ok=True)

    timestamp   = datetime.now().strftime("%d/%m/%Y %H:%M:%S")
    ts_filename = datetime.now().strftime("%Y%m%d_%H%M%S")

    if output_path is None:
        output_path = output_dir / f"report_{ts_filename}.xlsx"
    else:
        output_path = Path(output_path)

    # Buka existing untuk pertahankan riwayat, atau buat baru
    if output_path.exists():
        try:
            wb = openpyxl.load_workbook(output_path)
            logger.info("Membuka report existing: %s", output_path)
        except Exception as exc:
            logger.warning("Gagal buka existing (%s), buat baru.", exc)
            wb = Workbook()
    else:
        wb = Workbook()

    # Sheet DATA TERBARU — selalu dibuat ulang
    if "DATA TERBARU" in wb.sheetnames:
        del wb["DATA TERBARU"]
    ws_data = wb.create_sheet("DATA TERBARU", 0)
    _write_data_terbaru(ws_data, results)

    # Sheet RIWAYAT CEK — append saja
    if "RIWAYAT CEK" not in wb.sheetnames:
        ws_riwayat = wb.create_sheet("RIWAYAT CEK")
    else:
        ws_riwayat = wb["RIWAYAT CEK"]
    _append_riwayat(ws_riwayat, results, timestamp)

    # Bersihkan sheet default kosong
    if "Sheet" in wb.sheetnames and len(wb.sheetnames) > 1:
        del wb["Sheet"]

    wb.active = wb["DATA TERBARU"]
    wb.save(output_path)
    logger.info("Report disimpan: %s", output_path)
    return output_path
